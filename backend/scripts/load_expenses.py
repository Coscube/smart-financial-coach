#!/usr/bin/env python
"""
Script to load transaction data from an Excel file into the database.

The Excel file should have two sheets:
- "Expenses" with columns: Date and time | Category | Amount | Comment
- "Income" with columns: Date and time | Category | Amount | Comment

Row 1: Sheet title (ignored)
Row 2: Headers (ignored)
Row 3+: Data

Usage:
    poetry run python scripts/load_expenses.py <excel_file> [user_id]

Example:
    poetry run python scripts/load_expenses.py transactions.xlsx 1
"""

import sys
from pathlib import Path
from datetime import datetime, timezone
from decimal import Decimal
import asyncio

# Add backend directory to path
backend_dir = Path(__file__).parent.parent
sys.path.insert(0, str(backend_dir))

from openpyxl import load_workbook
from src.database import async_session_factory
from src.models.transaction import Transaction, TransactionType


def normalize_category(excel_category: str) -> str:
    """
    Normalize the Excel category string.
    
    Args:
        excel_category: Category string from Excel
        
    Returns:
        Normalized category string
    """
    if not excel_category:
        return "Uncategorized"
    return str(excel_category).strip()


async def load_transactions_from_excel(excel_file: str, user_id: int) -> None:
    """
    Load transactions from an Excel file into the database.
    
    Processes both "Expenses" and "Income" sheets.
    
    Args:
        excel_file: Path to the Excel file
        user_id: User ID to associate transactions with
    """
    
    # Check if file exists
    if not Path(excel_file).exists():
        print(f"Error: File '{excel_file}' not found")
        sys.exit(1)
    
    # Load workbook
    print(f"Loading workbook from {excel_file}...")
    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    print(f"Available sheets: {workbook.sheetnames}")
    
    DATE_COL = 0
    CATEGORY_COL = 1
    AMOUNT_COL = 2
    COMMENT_COL = 3
    
    all_transactions = []
    total_skipped = 0
    
    # Process each sheet
    sheets_config = [
        ("Expenses", TransactionType.EXPENSE),
        ("Income", TransactionType.INCOME),
    ]
    
    for sheet_name, transaction_type in sheets_config:
        if sheet_name not in workbook.sheetnames:
            print(f"Warning: Sheet '{sheet_name}' not found, skipping...")
            continue
        
        print(f"\n--- Processing {sheet_name} sheet ---")
        worksheet = workbook[sheet_name]
        
        transactions = []
        skipped = 0
        
        # Data starts from row 3 (row 1 is title, row 2 is headers)
        for row_idx, row in enumerate(worksheet.iter_rows(min_row=3, values_only=True), start=3):
            try:
                # Skip completely empty rows
                if not row or all(cell is None for cell in row):
                    continue
                
                # Extract values using correct column indices
                transaction_date = row[DATE_COL] if len(row) > DATE_COL else None
                category_str = row[CATEGORY_COL] if len(row) > CATEGORY_COL else None
                amount = row[AMOUNT_COL] if len(row) > AMOUNT_COL else None
                comment = row[COMMENT_COL] if len(row) > COMMENT_COL else None
                
                # Validate required fields (date and amount are required)
                if transaction_date is None or amount is None:
                    skipped += 1
                    continue
                
                # Convert amount to Decimal
                try:
                    amount = Decimal(str(amount))
                    if amount <= 0:
                        print(f"Warning: Row {row_idx} - invalid amount (must be positive): {amount}")
                        skipped += 1
                        continue
                except (ValueError, TypeError):
                    print(f"Warning: Row {row_idx} - invalid amount: {amount}")
                    skipped += 1
                    continue
                
                # Normalize category string
                category = normalize_category(category_str)
                
                # Handle date conversion
                if isinstance(transaction_date, datetime):
                    # Already a datetime object from openpyxl
                    if transaction_date.tzinfo is None:
                        transaction_date = transaction_date.replace(tzinfo=timezone.utc)
                elif isinstance(transaction_date, str):
                    try:
                        transaction_date = datetime.fromisoformat(transaction_date)
                    except ValueError:
                        try:
                            transaction_date = datetime.strptime(transaction_date, "%Y-%m-%d")
                        except ValueError:
                            print(f"Warning: Row {row_idx} - invalid date format: {transaction_date}")
                            skipped += 1
                            continue
                    if transaction_date.tzinfo is None:
                        transaction_date = transaction_date.replace(tzinfo=timezone.utc)
                else:
                    print(f"Warning: Row {row_idx} - unexpected date type: {type(transaction_date)}")
                    skipped += 1
                    continue
                
                # Clean up comment/description
                description = str(comment).strip() if comment else None
                
                # Create transaction object
                transaction = Transaction(
                    user_id=user_id,
                    transaction_type=transaction_type,
                    amount=amount,
                    category=category,
                    description=description,
                    transaction_date=transaction_date,
                )
                
                transactions.append(transaction)
                
                # Print sample rows
                if len(transactions) <= 3:
                    type_label = "EXPENSE" if transaction_type == TransactionType.EXPENSE else "INCOME"
                    print(f"  Row {row_idx}: [{type_label}] {amount} | {category} | {description or ''}")
            
            except Exception as e:
                print(f"Error processing row {row_idx}: {str(e)}")
                skipped += 1
                continue
        
        print(f"  Found {len(transactions)} valid {sheet_name.lower()} transactions")
        if skipped > 0:
            print(f"  Skipped {skipped} rows due to errors or missing data")
        
        all_transactions.extend(transactions)
        total_skipped += skipped
    
    # Insert all transactions into database
    if all_transactions:
        async with async_session_factory() as session:
            print(f"\nInserting {len(all_transactions)} total transactions into database...")
            session.add_all(all_transactions)
            await session.commit()
            print(f"Successfully inserted {len(all_transactions)} transactions!")
    else:
        print("\nNo valid transactions to insert")
    
    if total_skipped > 0:
        print(f"Total skipped rows: {total_skipped}")


async def main():
    """Main entry point."""
    if len(sys.argv) < 2:
        print("Usage: poetry run python scripts/load_expenses.py <excel_file> [user_id]")
        print("\nExample: poetry run python scripts/load_expenses.py transactions.xlsx 1")
        print("\nThe Excel file should have 'Expenses' and 'Income' sheets with columns:")
        print("  Date and time | Category | Amount | Comment")
        sys.exit(1)
    
    excel_file = sys.argv[1]
    user_id = int(sys.argv[2]) if len(sys.argv) > 2 else 1
    
    print(f"Loading transactions from {excel_file} for user {user_id}...\n")
    
    try:
        await load_transactions_from_excel(excel_file, user_id)
    except Exception as e:
        print(f"Error: {str(e)}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
